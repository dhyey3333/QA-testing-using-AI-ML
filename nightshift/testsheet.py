"""The test-case sheet QA teams keep: Excel first, the same rows as CSV, then into Jira.

The way most QA teams work: test cases are written in an Excel sheet, saved as CSV, imported into
Jira (or a test tool inside it), run by hand, and bugs are filed against them. Nightshift does the
same chain. After every run it writes that sheet, filled in, into the run folder:

    test-cases.xlsx   the sheet, formatted, with a coloured Status and links to Jira
    test-cases.csv    the same rows, for Jira's CSV import or any test-management tool
    test-cases.json   the same rows as data, so Nightshift can update the sheet later

Then `nightshift run --file-jira` (or "Send to Jira" in the web app) puts it into Jira: one issue per
test case, one per bug (jira.py, no duplicates), each bug linked to its test case, and the ticket keys
written back into the sheet.

Test data appears by name only ("email, password"), never its value. A cell that starts like a
formula (= + - @) gets a leading apostrophe, so text from a page can't run as one when opened.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
from pathlib import Path

import httpx

from .defects import analyse, load_results, record_issue
from .jira import JiraConfig, JiraError, _find_open, _issue_type, _post, file_jira_issues
from .result import RunResult
from .spec import Spec

SHEET = Path("test-cases.xlsx")  # inside the run folder; .csv and .json sit beside it
COLUMNS = ("Test Case ID", "Title", "Requirement", "Priority", "Preconditions", "Test Steps", "Test Data",
           "Expected Result", "Actual Result", "Status", "Bug ID", "Jira Test Case", "Executed On", "Executed By",
           "Evidence")
STATUS = {"pass": "Pass", "fail": "Fail", "flaky": "Flaky", "error": "Blocked"}  # error: the tester couldn't test it
WIDTHS = {"Test Case ID": 24, "Title": 36, "Requirement": 14, "Priority": 10, "Preconditions": 30, "Test Steps": 48,
          "Test Data": 18, "Expected Result": 42, "Actual Result": 48, "Status": 10, "Bug ID": 14,
          "Jira Test Case": 15, "Executed On": 17, "Executed By": 26, "Evidence": 28}
FILLS = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Flaky": "FFEB9C", "Blocked": "FFEB9C", "Not Run": "D9D9D9"}
TEST_TYPES = ("Test", "Test Case", "Task", "Story")  # a project without a Test issue type gets Tasks


def sheet_rows(specs: list[Spec], results: list[RunResult], base: Path) -> list[dict[str, str]]:
    """One row per test case, in the columns QA teams use, filled in from the run. Evidence links are
    relative to `base`, the folder the sheet goes in."""
    by_spec = {result.spec: result for result in results}
    rows = []
    for spec in specs:
        result = by_spec.get(spec.name)
        if spec.steps:
            steps = list(spec.steps)
        else:  # an API test: its requests are its steps
            steps = [f"{r.method} {r.path}" + (f" ({r.name})" if r.name else "") for r in spec.requests]
        precondition = f"A fresh browser at {spec.url}" if spec.kind == "ui" else f"The API at {spec.url}"
        if spec.session_from:
            precondition += f", logged in as in {spec.session_from}"
        evidence = ""
        if result is not None and (Path(result.out_dir) / "report.html").exists():
            try:
                evidence = Path(os.path.relpath(Path(result.out_dir) / "report.html", base)).as_posix()
            except ValueError:  # another drive: no relative link possible
                evidence = ""
        rows.append({
            "Test Case ID": spec.name,
            "Title": spec.title or (steps[0] if steps else spec.name),
            "Requirement": ", ".join(spec.requirements),
            "Priority": spec.priority,
            "Preconditions": precondition,
            "Test Steps": "\n".join(f"{i}. {step}" for i, step in enumerate(steps, 1)),
            "Test Data": ", ".join(spec.data),
            "Expected Result": "\n".join(f"• {item}" for item in spec.expect)
                               or ("• each request gets the answer its test expects" if spec.kind == "api" else ""),
            "Actual Result": _actual(result),
            "Status": STATUS.get(result.verdict, "Not Run") if result else "Not Run",
            "Bug ID": "",
            "Jira Test Case": "",
            "Executed On": result.started_at.replace("T", " ")[:16] if result else "",
            "Executed By": f"Nightshift ({result.model})" if result else "",
            "Evidence": evidence,
        })
    return rows


def sheet_for_old_run(run_dir: Path, results: list[RunResult], spec_files: list[Path]) -> Path:
    """The sheet for a run from before sheets existed, from its results and the test files as they are now
    (a test since deleted keeps its name and result, with its steps left blank)."""
    from dataclasses import replace

    from .spec import SpecError, load_spec

    specs: dict[str, Spec] = {}
    for path in spec_files:
        try:
            spec = load_spec(path, expand_env=False)  # never needs the secrets: only names reach the sheet
        except SpecError:
            continue
        specs[spec.name] = spec
    chosen = []
    for result in results:
        spec = specs.get(result.spec) or specs.get(result.spec.split("@")[0])  # "checkout@firefox" is checkout
        chosen.append(replace(spec, name=result.spec) if spec else Spec(name=result.spec, url=result.url, steps=(), expect=()))
    return write_sheet(sheet_rows(chosen, results, run_dir), run_dir / SHEET)


def _actual(result: RunResult | None) -> str:
    """What happened, in the words a tester would write in the Actual Result column."""
    if result is None:
        return ""
    if result.verdict == "pass":
        quotes = [quote for check in result.checks for quote in check.evidence][:3]
        text = "As expected." + (" Shown: " + "; ".join(f'"{q}"' for q in quotes) if quotes else "")
        loaded = [w for w in result.warnings if w.startswith("uncaught JS error while the page loaded")]
        return text + (f" Note: {loaded[0]}" if loaded else "")
    if result.verdict == "error":
        return f"Couldn't test: {result.reason}"
    if result.verdict == "flaky":
        return f"Failed, then passed when tried again: {result.reason}"
    return result.reason


def write_sheet(rows: list[dict[str, str]], path: Path, jira_url: str = "") -> Path:
    """The sheet as .xlsx, and the same rows as .csv and .json beside it. Returns the .xlsx path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Test cases"
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append([str(row.get(column, "") or "") for column in COLUMNS])
        line = sheet.max_row
        for index, column in enumerate(COLUMNS, 1):
            cell = sheet.cell(line, index)
            cell.data_type = "s"  # text, always: "=..." from a page is shown, never run as a formula
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            value = row.get(column, "")
            if column == "Status" and value in FILLS:
                cell.fill = PatternFill("solid", fgColor=FILLS[value])
                cell.font = Font(bold=True)
            elif column in ("Bug ID", "Jira Test Case") and value and jira_url:
                cell.hyperlink = f"{jira_url}/browse/{value.split(',')[0].strip()}"
                cell.style = "Hyperlink"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            elif column == "Evidence" and value:
                cell.hyperlink = value  # relative: opens from the downloaded run folder
                cell.style = "Hyperlink"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, column in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS[column]
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)

    with path.with_suffix(".csv").open("w", newline="", encoding="utf-8-sig") as handle:  # the BOM: Excel reads ₹
        writer = csv.writer(handle)
        writer.writerow(COLUMNS)
        writer.writerows([_safe(row.get(column, "")) for column in COLUMNS] for row in rows)
    path.with_suffix(".json").write_text(json.dumps({"jira_url": jira_url, "rows": rows}, indent=2, ensure_ascii=False),
                                         encoding="utf-8")
    return path


def _safe(value: str) -> str:
    """In a CSV, a cell starting with = + - @ (or a tab or return) is a formula to Excel. Text from a web
    page must never run as one, so it gets the apostrophe Excel itself uses to mean "this is text"."""
    value = str(value or "")
    return "'" + value if value[:1] in ("=", "+", "-", "@", "\t", "\r") else value


def write_markdown(rows: list[dict[str, str]], path: Path) -> Path:
    """The same test cases as a document to read."""
    lines = ["# Test cases", "", f"{len(rows)} case(s).", ""]
    for row in rows:
        lines += [f"## {row['Test Case ID']}: {row['Title']}", ""]
        facts = [f"{name}: {row[name]}" for name in ("Requirement", "Priority", "Status") if row.get(name)]
        lines += [" · ".join(facts), ""] if facts else []
        lines += [f"**Preconditions:** {row['Preconditions']}", "", "**Steps:**", "", row["Test Steps"], "",
                  "**Expected:**", "", row["Expected Result"], ""]
        if row.get("Test Data"):
            lines += [f"**Test data:** {row['Test Data']}", ""]
        if row.get("Actual Result"):
            lines += [f"**Actual result:** {row['Actual Result']}", ""]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def load_rows(run_dir: Path) -> tuple[list[dict[str, str]], str]:
    data = json.loads((run_dir / SHEET).with_suffix(".json").read_text(encoding="utf-8"))
    return data["rows"], data.get("jira_url", "")


def send_to_jira(run_dir: Path, config: JiraConfig, client: httpx.Client | None = None, log=None) -> dict:
    """The sheet into Jira: one issue per test case, one per bug, each bug linked to its test cases, and the
    keys written back into the sheet. Returns {"cases": [(test id, action, key)], "bugs": [(defect id, action, key)]}."""
    log = log or (lambda line: None)
    rows, _ = load_rows(run_dir)
    defects = analyse(load_results(run_dir), run_dir.parent)
    own = client is None
    client = client or config.client()
    try:
        cases = file_test_cases(rows, run_dir, config, client, log)
        case_keys = {test_id: key for test_id, _, key in cases}
        bugs = file_jira_issues(defects, run_dir, config, client, log) if defects else []
        by_id = {defect.id: defect for defect in defects}
        bug_keys: dict[str, list[str]] = {}
        for defect_id, action, key in bugs:
            record_issue(run_dir, defect_id, "jira", key, f"{config.url}/browse/{key}")
            for test_id in by_id[defect_id].specs:
                bug_keys.setdefault(test_id, []).append(key)
                if action == "created" and test_id in case_keys:
                    _link(client, case_keys[test_id], key, log)
    finally:
        if own:
            client.close()
    for row in rows:
        row["Jira Test Case"] = case_keys.get(row["Test Case ID"], row.get("Jira Test Case", ""))
        if row["Test Case ID"] in bug_keys:
            row["Bug ID"] = ", ".join(dict.fromkeys(bug_keys[row["Test Case ID"]]))
    write_sheet(rows, run_dir / SHEET, jira_url=config.url)
    return {"cases": cases, "bugs": bugs}


def file_test_cases(rows: list[dict[str, str]], run_dir: Path, config: JiraConfig, client: httpx.Client,
                    log) -> list[tuple[str, str, str]]:
    """One Jira issue per test case. A test case that has an issue already (found by its label, open or
    closed) gets this run's result as a comment instead; one already sent from this run is left alone."""
    done, issue_type = [], None
    for row in rows:
        test_id = row["Test Case ID"]
        if row.get("Jira Test Case"):
            done.append((test_id, "already sent", row["Jira Test Case"]))
            continue
        label = "nstc-" + hashlib.sha1(test_id.encode("utf-8")).hexdigest()[:10]
        existing = _find_open(client, config, label, open_only=False)
        if existing:
            _post(client, f"/rest/api/2/issue/{existing}/comment",
                  {"body": f"Run {run_dir.name}: *{row['Status']}*. {row['Actual Result']}"})
            done.append((test_id, "updated", existing))
            continue
        issue_type = issue_type or _issue_type(client, config, log, wanted=config.test_issue_type, fallbacks=TEST_TYPES)
        fields = {
            "project": {"key": config.project},
            "summary": f"Test case: {row['Title']}"[:250],
            "description": case_markup(row, run_dir.name),
            "issuetype": {"name": issue_type},
            "labels": ["nightshift", "test-case", label],
        }
        done.append((test_id, "created", _post(client, "/rest/api/2/issue", {"fields": fields})["key"]))
    return done


def case_markup(row: dict[str, str], run_name: str) -> str:
    """A test case in Jira's wiki markup: the sheet's row, laid out the way a person would write it."""
    lines = [f"*Test case ID:* {row['Test Case ID']}"]
    lines += [f"*{name}:* {row[name]}" for name in ("Requirement", "Priority") if row.get(name)]
    lines += [f"*Preconditions:* {row['Preconditions']}", "", "h3. Steps"]
    lines += ["# " + re.sub(r"^\d+\.\s*", "", step) for step in row["Test Steps"].splitlines()]
    lines += ["", "h3. Expected result"]
    lines += ["* " + item.removeprefix("• ") for item in row["Expected Result"].splitlines()]
    if row.get("Test Data"):
        lines += ["", f"*Test data:* {row['Test Data']}"]
    lines += ["", f"h3. Last run ({run_name})", f"*Status:* {row['Status']}", f"*Actual result:* {row['Actual Result']}"]
    if row.get("Executed On"):
        lines.append(f"*Executed on:* {row['Executed On']} by {row['Executed By']}")
    return "\n".join(lines)


def _link(client: httpx.Client, test_key: str, bug_key: str, log) -> None:
    """Link a bug to the test case that found it. A Jira without the usual "Relates" link type keeps
    both issues and says so; the link is a convenience, not worth failing the whole send over."""
    try:
        _post(client, "/rest/api/2/issueLink",
              {"type": {"name": "Relates"}, "inwardIssue": {"key": test_key}, "outwardIssue": {"key": bug_key}})
    except JiraError as exc:
        log(f"jira: could not link {bug_key} to {test_key} ({exc})")
