"""The test-case sheet: Excel first, the same rows as CSV, then into Jira (testsheet.py).

Runs use the scripted model against the demo shop; Jira is the stand-in server from test_jira.py.
"""

import csv
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from nightshift import cli
from nightshift.jira import JiraConfig
from nightshift.result import RunResult
from nightshift.spec import Spec
from nightshift.testsheet import COLUMNS, SHEET, load_rows, send_to_jira, sheet_rows, write_sheet
from test_e2e import CHECKOUT, CHECKOUT_EVIDENCE, LOGIN
from test_jira import jira  # noqa: F401 (the fixture)

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture
def finished_run(run, spec_for, tmp_path):
    """A run folder like `nightshift run` leaves: a passing login and a checkout that hits a server error."""
    results = [run("login", LOGIN, evidence=["Hi, Test Shopper"]),
               run("checkout", CHECKOUT, bugs={"checkout-500"}, evidence=CHECKOUT_EVIDENCE)]
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "summary.json").write_text(json.dumps([r.summary() for r in results]), encoding="utf-8")
    specs = [spec_for("login"), spec_for("checkout")]
    write_sheet(sheet_rows(specs, results, run_dir), run_dir / SHEET)
    return run_dir, specs


def test_a_run_becomes_an_excel_sheet_and_the_same_rows_as_csv(finished_run):
    run_dir, specs = finished_run
    sheet = load_workbook(run_dir / SHEET).active
    assert [c.value for c in sheet[1]] == list(COLUMNS) and sheet.freeze_panes == "B2"
    rows = {row[0].value: {c: cell.value for c, cell in zip(COLUMNS, row)} for row in sheet.iter_rows(min_row=2)}
    login, checkout = rows["login"], rows["checkout"]
    assert login["Status"] == "Pass" and login["Actual Result"].startswith('As expected. Shown: "Hi, Test Shopper"')
    assert login["Test Steps"] == "1. open the login page\n2. log in with the test account"
    assert login["Expected Result"] == "• the header greets the logged-in user by name"
    assert login["Test Data"] == "email, password" and login["Executed By"] == "Nightshift (scripted)"
    assert checkout["Status"] == "Fail" and "HTTP 500" in checkout["Actual Result"]
    assert sheet.cell(3, COLUMNS.index("Status") + 1).fill.fgColor.rgb.endswith("FFC7CE")  # red for a fail
    assert login["Evidence"].endswith("/report.html") and (run_dir / login["Evidence"]).exists()

    with (run_dir / SHEET).with_suffix(".csv").open(encoding="utf-8-sig") as handle:
        table = list(csv.DictReader(handle))
    assert [r["Test Case ID"] for r in table] == ["login", "checkout"]
    assert {k: v for k, v in table[0].items()} == {k: (v or "") for k, v in login.items()}
    # Test data by name only: the account's values are in no file.
    for path in [(run_dir / SHEET).with_suffix(s) for s in (".csv", ".json")]:
        assert "chai-time-42" not in path.read_text(encoding="utf-8")


def test_text_that_looks_like_a_formula_never_runs_as_one(tmp_path):
    spec = Spec(name="t", url="https://app.test/", steps=("s",), expect=("e",))
    result = RunResult(spec="t", url=spec.url, model="m", verdict="fail", out_dir=str(tmp_path),
                       reason='=HYPERLINK("http://evil.test","click")')
    write_sheet(sheet_rows([spec], [result], tmp_path), tmp_path / SHEET)
    cell = load_workbook(tmp_path / SHEET).active.cell(2, COLUMNS.index("Actual Result") + 1)
    assert cell.data_type == "s" and cell.value == '=HYPERLINK("http://evil.test","click")'  # shown, not run
    with (tmp_path / SHEET).with_suffix(".csv").open(encoding="utf-8-sig") as handle:
        assert next(csv.DictReader(handle))["Actual Result"].startswith("'=HYPERLINK")


def test_the_sheet_goes_into_jira_as_test_cases_and_bugs_linked_together(finished_run, jira):  # noqa: F811
    run_dir, _ = finished_run
    server = jira()
    config = JiraConfig(url=f"http://127.0.0.1:{server.server_port}", project="SHOP", email="qa@example.test", token="t")
    sent = send_to_jira(run_dir, config)
    # This project has no "Test" issue type, so test cases are Tasks; the bug is a Bug.
    assert [(t, a) for t, a, _ in sent["cases"]] == [("login", "created"), ("checkout", "created")]
    assert [a for _, a, _ in sent["bugs"]] == ["created"]
    cases = {t: k for t, _, k in sent["cases"]}
    bug = sent["bugs"][0][2]
    issue = server.issues[cases["login"]]["fields"]
    assert issue["issuetype"]["name"] == "Task" and issue["summary"] == "Test case: open the login page"
    assert "# open the login page" in issue["description"] and "*Status:* Pass" in issue["description"]
    assert server.issues[bug]["fields"]["issuetype"]["name"] == "Bug"
    assert server.links == [{"type": {"name": "Relates"}, "inwardIssue": {"key": cases["checkout"]},
                             "outwardIssue": {"key": bug}}]
    # The keys are written back into the sheet, with links to Jira.
    rows = {r["Test Case ID"]: r for r in load_rows(run_dir)[0]}
    assert rows["checkout"]["Bug ID"] == bug and rows["checkout"]["Jira Test Case"] == cases["checkout"]
    assert rows["login"]["Bug ID"] == ""
    link = load_workbook(run_dir / SHEET).active.cell(3, COLUMNS.index("Bug ID") + 1).hyperlink.target
    assert link == f"{config.url}/browse/{bug}"

    # Sent again: nothing is duplicated. The test cases are already in, the open bug gets a comment.
    again = send_to_jira(run_dir, config)
    assert [a for _, a, _ in again["cases"]] == ["already sent", "already sent"]
    assert [a for _, a, _ in again["bugs"]] == ["commented"] and len(server.issues) == 3


def test_a_later_run_adds_its_result_to_the_test_cases_already_in_jira(finished_run, jira, tmp_path):  # noqa: F811
    run_dir, specs = finished_run
    server = jira()
    config = JiraConfig(url=f"http://127.0.0.1:{server.server_port}", project="SHOP", email="qa@example.test", token="t")
    send_to_jira(run_dir, config)
    later = tmp_path / "later"
    later.mkdir()
    (later / "summary.json").write_text((run_dir / "summary.json").read_text(encoding="utf-8"), encoding="utf-8")
    rows, _ = load_rows(run_dir)
    write_sheet([{**row, "Bug ID": "", "Jira Test Case": ""} for row in rows], later / SHEET)
    sent = send_to_jira(later, config)
    assert [a for _, a, _ in sent["cases"]] == ["updated", "updated"]
    assert any(body.startswith("Run later: *Pass*") for _, body in server.comments)


def test_every_run_writes_the_sheet(shop, base_url, tmp_path):
    # Its own process: this one already has a browser open for the other tests.
    env = {**os.environ, "MODEL_BASE_URL": "http://127.0.0.1:9/v1"}  # replayed from saved paths: no model
    done = subprocess.run([sys.executable, "-m", "nightshift", "run", str(ROOT / "specs" / "login.yaml"),
                           "--recordings", str(ROOT / "ci" / "recordings"), "--no-record", "--retries", "0",
                           "--base-url", base_url, "--out", str(tmp_path), "--quiet"],
                          env=env, capture_output=True, text=True, encoding="utf-8", timeout=300)
    assert done.returncode == 0, done.stdout + done.stderr
    run_dir = next(tmp_path.glob("*/summary.json")).parent
    assert (run_dir / "test-cases.xlsx").exists() and (run_dir / "test-cases.csv").exists()
    assert [r["Status"] for r in load_rows(run_dir)[0]] == ["Pass"]
    assert "test cases: " in done.stdout


def test_the_cases_command_writes_the_sheet_before_anything_runs(tmp_path, capsys):
    out = tmp_path / "test-cases.csv"
    assert cli.main(["cases", str(ROOT / "specs" / "login.yaml"), "--to", str(out)]) == 0
    assert out.exists() and out.with_suffix(".xlsx").exists() and out.with_suffix(".md").exists()
    assert [r["Status"] for r in load_rows(tmp_path)[0]] == ["Not Run"]
